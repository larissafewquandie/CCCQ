#!/usr/bin/env python3
"""Turn a pasted "Activities and Clients" search result into an .xlsx.

The case management system copies each result as a vertical block of lines:

    HOLMES-000000          <- activity ID (SURNAME-number)
    DV                     <- program
    Jane HOLMES            <- client
    17 Apr 2026 (Fri)      <- commencement
    22 Apr 2026 (Wed)      <- last session (may carry a footnote digit, or "No sessions")
    Worker Name            <- worker
    Rockhampton            <- location
    Closed                 <- status

Usage: extract_activities.py INPUT.txt OUTPUT.xlsx [--match NAME]
"""
import argparse
import re
import sys
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ID_RE = re.compile(r"^[A-Z][A-Z'/\- ]*-\d{4,}$")
DATE_RE = re.compile(r"^(\d{1,2} [A-Z][a-z]{2} \d{4})(?: \([A-Za-z]{3}\))?\s*(\d*)$")
HEADER_RE = re.compile(r"^Searching for (?P<what>.+?) where (?P<where>.+?)\.\s*(?:Sorting by (?P<sort>.+?)\.)?\s*$")
FIELDS = ["Activity ID", "Program", "Client", "Commencement", "Last Session",
          "Worker", "Location", "Status"]


def parse_date(text):
    m = DATE_RE.match(text.strip())
    if not m:
        return None, ""
    return datetime.strptime(m.group(1), "%d %b %Y").date(), m.group(2)


def parse(lines):
    meta, records, current = {}, [], None
    for raw in lines:
        line = raw.strip()
        if not line or line == "|":
            continue
        header = HEADER_RE.match(line)
        if header:
            meta = {k: v for k, v in header.groupdict().items() if v}
            continue
        if ID_RE.match(line):
            current = [line]
            records.append(current)
        elif current is not None:
            current.append(line)
    return meta, [build(r) for r in records]


def build(values):
    rec = dict(zip(FIELDS, values + [""] * (len(FIELDS) - len(values))))
    issues = []
    if len(values) < len(FIELDS):
        issues.append(f"Incomplete record ({len(values)}/{len(FIELDS)} fields) - paste may be cut off")
    elif len(values) > len(FIELDS):
        issues.append("Extra lines: " + " | ".join(values[len(FIELDS):]))

    commenced, _ = parse_date(rec["Commencement"])
    if rec["Commencement"] and commenced is None:
        issues.append(f"Unreadable commencement: {rec['Commencement']}")
    rec["Commencement"] = commenced

    last_raw = rec["Last Session"]
    last, footnote = parse_date(last_raw)
    if last_raw and last is None and last_raw.lower() != "no sessions":
        issues.append(f"Unreadable last session: {last_raw}")
    rec["Last Session"] = last
    rec["No Sessions"] = "Yes" if last_raw.lower() == "no sessions" else ""
    rec["Footnote"] = footnote
    rec["Days Open"] = (last - commenced).days if commenced and last else None
    rec["Issues"] = "; ".join(issues)
    return rec


def write_xlsx(meta, records, path, match):
    wb = Workbook()
    ws = wb.active
    ws.title = "Activities"
    columns = FIELDS + ["Days Open", "No Sessions", "Footnote"]
    if match:
        columns.append("First Name Match")
    columns.append("Issues")
    ws.append(columns)

    flag = PatternFill("solid", fgColor="FFF2CC")
    for rec in records:
        if match:
            first = rec["Client"].split(" ")[0] if rec["Client"] else ""
            rec["First Name Match"] = "Exact" if first.lower() == match.lower() else "Variant"
        ws.append([rec.get(c) for c in columns])
        if rec["Issues"]:
            for cell in ws[ws.max_row]:
                cell.fill = flag

    for col in ("Commencement", "Last Session"):
        letter = get_column_letter(columns.index(col) + 1)
        for cell in ws[letter][1:]:
            cell.number_format = "DD MMM YYYY"
    if records:
        ref = f"A1:{get_column_letter(len(columns))}{len(records) + 1}"
        table = Table(displayName="Activities", ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
    ws.freeze_panes = "A2"
    for i, col in enumerate(columns, 1):
        width = max([len(col)] + [len(str(ws.cell(r, i).value or "")) for r in range(2, ws.max_row + 1)])
        ws.column_dimensions[get_column_letter(i)].width = min(width + 2, 60)

    summary = wb.create_sheet("Summary")
    rows = [("Search", meta.get("what", "")), ("Criteria", meta.get("where", "")),
            ("Sort", meta.get("sort", "")), ("Extracted", datetime.now().strftime("%d %b %Y %H:%M")),
            ("Records", len(records)), ("Records with issues", sum(1 for r in records if r["Issues"]))]
    for label, key in (("By status", "Status"), ("By location", "Location"), ("By worker", "Worker")):
        rows.append(("", ""))
        rows.append((label, ""))
        counts = {}
        for r in records:
            counts[r[key] or "(blank)"] = counts.get(r[key] or "(blank)", 0) + 1
        rows.extend(sorted(counts.items(), key=lambda kv: (-kv[1], kv[0])))
    for row in rows:
        summary.append(row)
        if row[0] in ("By status", "By location", "By worker"):
            summary.cell(summary.max_row, 1).font = Font(bold=True)
    summary.column_dimensions["A"].width = 30
    summary.column_dimensions["B"].width = 40
    wb.save(path)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("input", help="text file holding the pasted search results ('-' for stdin)")
    ap.add_argument("output", help="path of the .xlsx to write")
    ap.add_argument("--match", help="searched first name; flags exact vs variant matches (e.g. Kristy vs Kristyn)")
    args = ap.parse_args()

    text = sys.stdin.read() if args.input == "-" else open(args.input, encoding="utf-8").read()
    meta, records = parse(text.splitlines())
    if not records:
        sys.exit("No records found - expected lines like SURNAME-123456 starting each result.")
    write_xlsx(meta, records, args.output, args.match)

    print(f"{len(records)} records -> {args.output}")
    for r in records:
        if r["Issues"]:
            print(f"  ! {r['Activity ID']}: {r['Issues']}")


if __name__ == "__main__":
    main()
